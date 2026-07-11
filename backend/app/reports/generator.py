from __future__ import annotations

import csv
from copy import deepcopy
from datetime import datetime
import hashlib
import json
import os
from pathlib import Path
import re
from typing import Any
import uuid
from xml.sax.saxutils import escape

from app.core.paths import ASSETS_DIR, REPORTS_DIR
from app.core.time import utc_now

from .templates import REPORT_TITLE


SAFE_NAME = re.compile(r"[^A-Za-z0-9_.-]+")
DOCX_FONT_NAME = "Noto Sans CJK SC"


class ReportGenerator:
    def __init__(self, output_dir: Path = REPORTS_DIR) -> None:
        self.output_dir = Path(output_dir)
        self.pdf_font_name, self.pdf_font_path, self.pdf_font_embedded = self._register_pdf_font()

    def generate(
        self,
        session: dict,
        steps: list[dict],
        *,
        metadata: dict | None = None,
    ) -> dict:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        report_id = uuid.uuid4().hex[:8]
        generated_at = utc_now()
        result = str(session.get("overall_result") or "UNKNOWN").upper()
        base = self.output_dir / self._base_name(session, result, report_id, generated_at)
        payload = self._payload(session, steps, metadata or {}, generated_at)
        files: dict[str, str] = {}
        statuses: dict[str, str] = {}
        errors: dict[str, str] = {}

        writers = {
            "json": lambda path: path.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2, default=str),
                encoding="utf-8",
            ),
            "docx": lambda path: self._write_docx(path, payload),
            "pdf": lambda path: self._write_pdf(path, payload),
            "csv": lambda path: self._write_csv(path, payload),
        }
        for report_type, writer in writers.items():
            path = base.with_suffix(f".{report_type}")
            try:
                writer(path)
                if not path.is_file() or path.stat().st_size == 0:
                    raise RuntimeError("generated file is empty")
            except Exception as exc:
                statuses[report_type] = f"{report_type.upper()}_FAILED"
                errors[report_type] = f"{type(exc).__name__}: {exc}"
                if path.exists():
                    path.unlink()
            else:
                files[report_type] = str(path)
                statuses[report_type] = "COMPLETED"

        if "json" not in files:
            raise RuntimeError(f"JSON report generation failed: {errors.get('json')}")
        return {
            "id": report_id,
            "session_id": session["id"],
            "result": result,
            "generated_at": generated_at,
            "files": files,
            "statuses": statuses,
            "errors": errors,
            "font": {
                "name": self.pdf_font_name,
                "path": str(self.pdf_font_path) if self.pdf_font_path else None,
                "embedded": self.pdf_font_embedded,
            },
            "hashes": {
                kind: hashlib.sha256(Path(path).read_bytes()).hexdigest()
                for kind, path in files.items()
            },
        }

    @staticmethod
    def _base_name(
        session: dict[str, Any], result: str, report_id: str, generated_at: str
    ) -> str:
        timestamp = generated_at.replace("-", "").replace(":", "").replace("T", "_")[:15]
        chassis = SAFE_NAME.sub("-", str(session.get("chassis_no") or "unknown"))[:40]
        vin_last = SAFE_NAME.sub("", str(session.get("vin") or ""))[-6:] or "novin"
        session_id = SAFE_NAME.sub("-", str(session.get("id") or "session"))[:48]
        return f"{timestamp}_{chassis}_{vin_last}_{result}_{session_id}_{report_id}"

    @staticmethod
    def _payload(
        session: dict[str, Any],
        steps: list[dict[str, Any]],
        metadata: dict[str, Any],
        generated_at: str,
    ) -> dict[str, Any]:
        normalized_steps = deepcopy(steps)
        assertions = [
            {
                **assertion,
                "step_id": step.get("id"),
                "step_name": step.get("name"),
                "step_order": step.get("order"),
            }
            for step in normalized_steps
            for assertion in step.get("assertions", [])
        ]
        passed = sum(step.get("result") == "PASS" for step in normalized_steps)
        failed = sum(step.get("result") == "FAIL" for step in normalized_steps)
        traceability = {
            "software_version": session.get("software_version") or metadata.get("software_version"),
            "dbc_hash": session.get("dbc_hash") or metadata.get("dbc_hash"),
            "config_hash": session.get("config_hash") or metadata.get("config_hash"),
            "config_profile": metadata.get("config_profile"),
            "test_plan_id": session.get("plan_id") or session.get("test_plan_id") or metadata.get("test_plan_id"),
            "test_plan_version": session.get("plan_version") or metadata.get("test_plan_version"),
            "operator": session.get("operator"),
            **metadata,
        }
        return {
            "schema_version": "2.0",
            "title": REPORT_TITLE,
            "generated_at": generated_at,
            "session": deepcopy(session),
            "summary": {
                "result": session.get("overall_result") or "UNKNOWN",
                "status": session.get("status"),
                "step_total": len(normalized_steps),
                "step_passed": passed,
                "step_failed": failed,
                "assertion_total": len(assertions),
                "assertion_failed": sum(item.get("result") == "FAIL" for item in assertions),
                "failure_reason": session.get("failure_reason") or "",
                "safe_stop": session.get("safe_stop_result"),
            },
            "traceability": traceability,
            "steps": normalized_steps,
            "assertions": assertions,
        }

    def _register_pdf_font(self) -> tuple[str, Path | None, bool]:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.pdfbase.ttfonts import TTFont

        configured = os.getenv("CHASSIS_REPORT_FONT", "").strip()
        candidates = [
            Path(configured) if configured else None,
            # The packaged TrueType font keeps Chinese text portable and preserves
            # Unicode extraction in both preview and verification tools.
            ASSETS_DIR / "fonts" / "NotoSansSC-VF.ttf",
            ASSETS_DIR / "fonts" / "NotoSansCJKsc-Regular.otf",
            Path(r"C:\Windows\Fonts\simhei.ttf"),
            Path(r"C:\Windows\Fonts\msyh.ttc"),
            Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
            Path("/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc"),
        ]
        for candidate in candidates:
            if candidate and candidate.is_file():
                try:
                    pdfmetrics.registerFont(TTFont("ChassisCJK", str(candidate)))
                    return "ChassisCJK", candidate, True
                except Exception:
                    continue
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        return "STSong-Light", None, False

    @staticmethod
    def _set_docx_font(run, name: str = DOCX_FONT_NAME, size: int | None = None) -> None:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
        if size:
            run.font.size = Pt(size)

    def _write_docx(self, path: Path, payload: dict[str, Any]) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Pt

        document = Document()
        normal = document.styles["Normal"]
        normal.font.name = DOCX_FONT_NAME
        normal.font.size = Pt(9)
        title = document.add_heading(payload["title"], 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            self._set_docx_font(run, size=20)

        session = payload["session"]
        trace = payload["traceability"]
        summary = payload["summary"]
        document.add_heading("1. 检测对象与结论", level=1)
        self._docx_table(
            document,
            [
                ("会话 ID", session.get("id")),
                ("底盘编号", session.get("chassis_no")),
                ("VIN", session.get("vin")),
                ("序列号", session.get("serial_no")),
                ("工位", session.get("station_id")),
                ("操作员", session.get("operator")),
                ("开始时间", session.get("started_at")),
                ("结束时间", session.get("ended_at")),
                ("检测结论", summary.get("result")),
                ("失败原因", summary.get("failure_reason") or "-"),
            ],
        )
        document.add_heading("2. 追溯信息", level=1)
        self._docx_table(document, [(str(key), value) for key, value in trace.items()])
        document.add_heading("3. 检测步骤", level=1)
        step_table = document.add_table(rows=1, cols=6)
        step_table.style = "Table Grid"
        for cell, text in zip(step_table.rows[0].cells, ["序号", "步骤", "状态", "结果", "耗时(ms)", "失败原因"]):
            cell.text = text
        for step in payload["steps"]:
            cells = step_table.add_row().cells
            values = [step.get("order"), step.get("name"), step.get("status"), step.get("result"), step.get("duration_ms"), step.get("failure_reason") or "-"]
            for cell, value in zip(cells, values):
                cell.text = str(value if value is not None else "-")
        document.add_heading("4. 检测断言", level=1)
        assertion_table = document.add_table(rows=1, cols=8)
        assertion_table.style = "Table Grid"
        headers = ["步骤", "断言", "信号", "阈值", "测量值", "单位", "质量/来源", "结果"]
        for cell, text in zip(assertion_table.rows[0].cells, headers):
            cell.text = text
        for assertion in payload["assertions"]:
            cells = assertion_table.add_row().cells
            values = [
                assertion.get("step_name"),
                assertion.get("description") or assertion.get("assertion_id"),
                assertion.get("signal_name"),
                json.dumps(assertion.get("threshold"), ensure_ascii=False, default=str),
                json.dumps(assertion.get("measured_value"), ensure_ascii=False, default=str),
                assertion.get("unit"),
                f'{assertion.get("quality", "-")} / {assertion.get("source_can_id", "-")}',
                assertion.get("result"),
            ]
            for cell, value in zip(cells, values):
                cell.text = str(value if value not in (None, "") else "-")
        document.add_heading("5. 结论与签名", level=1)
        document.add_paragraph(
            f'检测结论：{summary["result"]}；通过步骤 {summary["step_passed"]}/{summary["step_total"]}；'
            f'失败断言 {summary["assertion_failed"]}/{summary["assertion_total"]}。'
        )
        document.add_paragraph("操作员签名：________________    质量工程师：________________")
        for paragraph in document.paragraphs:
            for run in paragraph.runs:
                self._set_docx_font(run)
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            self._set_docx_font(run, size=8)
        document.save(path)

    def _docx_table(self, document, rows: list[tuple[str, Any]]) -> None:
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key, value in rows:
            cells = table.add_row().cells
            cells[0].text = str(key)
            cells[1].text = str(value if value not in (None, "") else "-")

    def _write_pdf(self, path: Path, payload: dict[str, Any]) -> None:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

        styles = getSampleStyleSheet()
        body = ParagraphStyle("CJKBody", parent=styles["BodyText"], fontName=self.pdf_font_name, fontSize=8.2, leading=12, wordWrap="CJK")
        heading = ParagraphStyle("CJKHeading", parent=styles["Heading2"], fontName=self.pdf_font_name, fontSize=12, leading=16, textColor=colors.HexColor("#17365D"), spaceBefore=8, spaceAfter=5)
        title = ParagraphStyle("CJKTitle", parent=styles["Title"], fontName=self.pdf_font_name, fontSize=20, leading=26, alignment=TA_CENTER, textColor=colors.HexColor("#0B3D70"))
        document = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title=payload["title"], author=str(payload["session"].get("operator") or ""))
        story = [Paragraph(escape(payload["title"]), title), Spacer(1, 5 * mm)]
        session = payload["session"]
        summary = payload["summary"]
        story.append(Paragraph("1. 检测对象与结论", heading))
        info = [
            ["会话 ID", session.get("id"), "检测结论", summary.get("result")],
            ["底盘编号", session.get("chassis_no"), "VIN", session.get("vin")],
            ["序列号", session.get("serial_no"), "工位", session.get("station_id")],
            ["操作员", session.get("operator"), "检测方案", payload["traceability"].get("test_plan_id")],
            ["开始时间", session.get("started_at"), "结束时间", session.get("ended_at")],
            ["失败原因", summary.get("failure_reason") or "-", "安全停车", json.dumps(summary.get("safe_stop"), ensure_ascii=False, default=str) if summary.get("safe_stop") else "-"],
        ]
        story.append(self._pdf_table(info, body, [28 * mm, 62 * mm, 28 * mm, 62 * mm]))
        story.append(Paragraph("2. 追溯信息", heading))
        trace_rows = [[key, value if value not in (None, "") else "-"] for key, value in payload["traceability"].items()]
        story.append(self._pdf_table(trace_rows, body, [45 * mm, 135 * mm]))
        story.append(Paragraph("3. 检测步骤", heading))
        step_rows = [["序号", "步骤", "状态", "结果", "耗时(ms)", "失败原因"]]
        step_rows.extend(
            [step.get("order"), step.get("name"), step.get("status"), step.get("result"), step.get("duration_ms") or "-", step.get("failure_reason") or "-"]
            for step in payload["steps"]
        )
        story.append(self._pdf_table(step_rows, body, [12 * mm, 38 * mm, 23 * mm, 20 * mm, 20 * mm, 67 * mm], header=True))
        story.append(PageBreak())
        story.append(Paragraph("4. 检测断言", heading))
        assertion_rows = [["步骤", "断言", "信号", "阈值", "测量值", "来源", "结果"]]
        for item in payload["assertions"]:
            assertion_rows.append(
                [
                    item.get("step_name"),
                    item.get("description") or item.get("assertion_id"),
                    item.get("signal_name"),
                    json.dumps(item.get("threshold"), ensure_ascii=False, default=str),
                    json.dumps(item.get("measured_value"), ensure_ascii=False, default=str),
                    f'{item.get("quality", "-")}/{item.get("source_can_id", "-")}',
                    item.get("result"),
                ]
            )
        story.append(self._pdf_table(assertion_rows, body, [25 * mm, 42 * mm, 28 * mm, 28 * mm, 25 * mm, 20 * mm, 14 * mm], header=True))
        if any(
            "完整值见 JSON/CSV 报告" in self._pdf_cell_value(item.get("measured_value"))
            for item in payload["assertions"]
        ):
            story.append(Spacer(1, 3 * mm))
            story.append(Paragraph("注：完整值见 JSON/CSV 报告。", body))
        story.append(Paragraph("5. 结论与签名", heading))
        conclusion = (
            f'检测结论：{summary["result"]}；通过步骤 {summary["step_passed"]}/{summary["step_total"]}；'
            f'失败断言 {summary["assertion_failed"]}/{summary["assertion_total"]}。'
        )
        story.extend([Paragraph(escape(conclusion), body), Spacer(1, 8 * mm), Paragraph("操作员签名：________________　质量工程师：________________", body)])

        def page_footer(canvas, doc):
            canvas.saveState()
            canvas.setFont(self.pdf_font_name, 7)
            canvas.setFillColor(colors.HexColor("#5B6B7C"))
            canvas.drawString(14 * mm, 8 * mm, str(session.get("id") or ""))
            canvas.drawRightString(A4[0] - 14 * mm, 8 * mm, f"第 {doc.page} 页")
            canvas.restoreState()

        document.build(story, onFirstPage=page_footer, onLaterPages=page_footer)

    def _pdf_table(self, rows, body_style, widths, header: bool = False):
        from reportlab.lib import colors
        from reportlab.platypus import Paragraph, Table, TableStyle

        converted = [
            [Paragraph(escape(self._pdf_cell_value(value)), body_style) for value in row]
            for row in rows
        ]
        table = Table(converted, colWidths=widths, repeatRows=1 if header else 0)
        commands = [
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#8EA9C1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if header:
            commands.extend(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#17365D")),
                ]
            )
        table.setStyle(TableStyle(commands))
        return table

    @staticmethod
    def _pdf_cell_value(value: Any, limit: int = 480) -> str:
        if value in (None, ""):
            return "-"
        if isinstance(value, (dict, list, tuple)):
            text = json.dumps(value, ensure_ascii=False, default=str)
        else:
            text = str(value)
        if len(text) <= limit:
            return text
        omitted = len(text) - limit
        return f"{text[:limit]}… [已省略 {omitted} 字符，完整值见 JSON/CSV 报告]"

    @staticmethod
    def _write_csv(path: Path, payload: dict[str, Any]) -> None:
        trace = payload["traceability"]
        session = payload["session"]
        columns = [
            "session_id",
            "chassis_no",
            "vin",
            "operator",
            "result",
            "software_version",
            "dbc_hash",
            "config_hash",
            "test_plan_id",
            "test_plan_version",
            "step_order",
            "step_id",
            "step_name",
            "step_result",
            "assertion_id",
            "description",
            "signal_name",
            "threshold",
            "measured_value",
            "unit",
            "quality",
            "source_can_id",
            "source_channel",
            "assertion_result",
            "failure_reason",
        ]
        with path.open("w", newline="", encoding="utf-8-sig") as stream:
            writer = csv.DictWriter(stream, fieldnames=columns)
            writer.writeheader()
            for step in payload["steps"]:
                assertions = step.get("assertions") or [{}]
                for assertion in assertions:
                    writer.writerow(
                        {
                            "session_id": session.get("id"),
                            "chassis_no": session.get("chassis_no"),
                            "vin": session.get("vin"),
                            "operator": session.get("operator"),
                            "result": payload["summary"]["result"],
                            "software_version": trace.get("software_version"),
                            "dbc_hash": trace.get("dbc_hash"),
                            "config_hash": trace.get("config_hash"),
                            "test_plan_id": trace.get("test_plan_id"),
                            "test_plan_version": trace.get("test_plan_version"),
                            "step_order": step.get("order"),
                            "step_id": step.get("id"),
                            "step_name": step.get("name"),
                            "step_result": step.get("result"),
                            "assertion_id": assertion.get("assertion_id") or assertion.get("id"),
                            "description": assertion.get("description"),
                            "signal_name": assertion.get("signal_name"),
                            "threshold": json.dumps(assertion.get("threshold"), ensure_ascii=False, default=str),
                            "measured_value": json.dumps(assertion.get("measured_value"), ensure_ascii=False, default=str),
                            "unit": assertion.get("unit"),
                            "quality": assertion.get("quality"),
                            "source_can_id": assertion.get("source_can_id"),
                            "source_channel": assertion.get("source_channel"),
                            "assertion_result": assertion.get("result"),
                            "failure_reason": assertion.get("failure_reason") or step.get("failure_reason"),
                        }
                    )
