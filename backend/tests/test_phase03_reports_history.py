from __future__ import annotations

import csv
from contextlib import contextmanager
from datetime import datetime, timezone
import json
from pathlib import Path

import fitz
from docx import Document
from fastapi.testclient import TestClient

from app.eol.models import AssertionOutcome
from app.main import app
from app.reports.generator import ReportGenerator
from app.reports.service import ReportService
from app.services.app_state import state
from app.services.history_service import HistoryService
from app.storage.database import Database
from app.storage.repositories import Repositories
from app.storage.uow import EolUnitOfWork


def _utc(hour: int, minute: int = 0) -> str:
    return datetime(2026, 4, 1, hour, minute, tzinfo=timezone.utc).isoformat()


@contextmanager
def isolated_data_services(tmp_path: Path):
    saved = {
        "database": state.database,
        "repositories": state.repositories,
        "reports": state.reports,
        "report_service": state.report_service,
        "history_service": state.history_service,
        "eol_uow": state.eol_uow,
    }
    database = Database(tmp_path / "phase03.sqlite")
    generator = ReportGenerator(tmp_path / "reports")
    state.database = database
    state.repositories = Repositories(database)
    state.reports = generator
    state.eol_uow = EolUnitOfWork(database)
    state.report_service = ReportService(state)
    state.history_service = HistoryService(state)
    state.history_service.export_root = tmp_path / "history-exports"
    state.history_service.export_root.mkdir(parents=True, exist_ok=True)
    try:
        yield database, generator
    finally:
        state.database = saved["database"]
        state.repositories = saved["repositories"]
        state.reports = saved["reports"]
        state.report_service = saved["report_service"]
        state.history_service = saved["history_service"]
        state.eol_uow = saved["eol_uow"]
        database.close()


def seed_session(database: Database, generator: ReportGenerator, session_id: str, result: str) -> dict:
    is_fail = result == "FAIL"
    database.execute(
        "INSERT INTO test_sessions(id,chassis_no,vin,serial_no,operator,station_id,test_plan_id,plan_version,status,overall_result,started_at,ended_at,remarks,dbc_hash,config_hash,software_version,failure_reason) "
        "VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (
            session_id,
            f"YL-{result}-001",
            f"L{result}000000000001",
            f"SN-{result}-001",
            "op-phase03",
            "EOL-STATION-01",
            "standard-eol",
            "1.0.2",
            "FAILED" if is_fail else "PASSED",
            result,
            _utc(10),
            _utc(10, 12),
            "阶段三中文报告验证",
            "dbc-phase03-hash",
            "config-phase03-hash",
            "v1.0.2",
            "制动反馈缺失" if is_fail else "",
        ),
    )
    steps = []
    for order, (step_id, name) in enumerate((("POWER_ON", "上电自检"), ("BMS_CHECK", "BMS检测"), ("BRAKE_STOP", "制动停止检测")), 1):
        step_result = "FAIL" if is_fail and step_id == "BRAKE_STOP" else "PASS"
        cursor = database.execute(
            "INSERT INTO test_steps(session_id,step_id,step_order,name,status,result,started_at,ended_at,duration_ms,failure_reason,command_summary,measurement_summary) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                session_id,
                step_id,
                order,
                name,
                "FAILED" if step_result == "FAIL" else "PASSED",
                step_result,
                _utc(10, order),
                _utc(10, order + 1),
                1000,
                "制动反馈缺失" if step_result == "FAIL" else "",
                json.dumps({"can_id": "0x121"}, ensure_ascii=False),
                json.dumps({"BMS总压": 329.6}, ensure_ascii=False),
            ),
        )
        assertion = AssertionOutcome(
            assertion_id=f"{step_id}-ASSERT",
            description=f"{name}断言",
            signal_name="Brake_Status" if step_id == "BRAKE_STOP" else "BMS_Voltage",
            operator="eq" if step_id == "BRAKE_STOP" else "between",
            threshold=True if step_id == "BRAKE_STOP" else [280, 360],
            measured_value=False if step_result == "FAIL" else (True if step_id == "BRAKE_STOP" else 329.6),
            unit="bool" if step_id == "BRAKE_STOP" else "V",
            result=step_result,
            severity="failure",
            failure_reason="制动反馈缺失" if step_result == "FAIL" else "",
            quality="good",
            source_can_id="0x51" if step_id == "BRAKE_STOP" else "0x100",
            source_channel="CAN2" if step_id == "BRAKE_STOP" else "CAN1",
            source_timestamp=_utc(10, order),
        )
        database.execute(
            "INSERT INTO test_assertions(session_id,step_id,assertion_id,description,signal_name,operator,threshold_json,measured_value,unit,result,severity,failure_reason,quality,source_can_id,source_channel,source_timestamp) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (
                session_id,
                step_id,
                assertion.assertion_id,
                assertion.description,
                assertion.signal_name,
                assertion.operator,
                json.dumps(assertion.threshold, ensure_ascii=False),
                json.dumps(assertion.measured_value, ensure_ascii=False),
                assertion.unit,
                assertion.result,
                assertion.severity,
                assertion.failure_reason,
                assertion.quality,
                assertion.source_can_id,
                assertion.source_channel,
                assertion.source_timestamp,
            ),
        )
        steps.append(
            {
                "id": step_id,
                "order": order,
                "name": name,
                "status": "FAILED" if step_result == "FAIL" else "PASSED",
                "result": step_result,
                "duration_ms": 1000,
                "failure_reason": assertion.failure_reason,
                "assertions": [assertion.model_dump()],
            }
        )

    database.execute(
        "INSERT INTO raw_can_frames(session_id,timestamp_utc,channel,direction,can_id_hex,is_extended,is_remote,dlc,data_hex,parse_status,message_name) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
        (session_id, _utc(10, 5), "CAN1", "RX", "0x100", 0, 0, 8, "0C E0 FF F4 56 00 00 00", "ok", "BMS_Status"),
    )
    for minute, value in ((5, 329.6), (6, 330.1)):
        database.execute(
            "INSERT INTO decoded_signals(session_id,timestamp_utc,channel,can_id_hex,message_name,signal_name,raw_value,physical_value,unit,quality,threshold_status) VALUES (?,?,?,?,?,?,?,?,?,?,?)",
            (session_id, _utc(10, minute), "CAN1", "0x100", "BMS_Status", "BMS_Voltage", str(int(value * 10)), str(value), "V", "good", "normal"),
        )
    if is_fail:
        database.execute(
            "INSERT INTO alarms(session_id,timestamp_utc,channel,can_id_hex,signal_name,level,level_label,status,description,related_step_id) VALUES (?,?,?,?,?,?,?,?,?,?)",
            (session_id, _utc(10, 9), "CAN2", "0x77", "Brake_Error", 2, "Fault", "ACTIVE", "制动故障", "BRAKE_STOP"),
        )
    state.repositories.action("create_phase03_session", session_id, {"result": result}, operator="op-phase03", trace_id=f"trace-{result.lower()}", session_id=session_id, role="operator")

    session = state.repositories.session(session_id)
    generated = generator.generate(
        session,
        steps,
        metadata={
            "software_version": "v1.0.2",
            "dbc_hash": "dbc-phase03-hash",
            "config_hash": "config-phase03-hash",
            "test_plan_id": "standard-eol",
            "test_plan_version": "1.0.2",
            "operator": "op-phase03",
        },
    )
    generated["database_ids"] = state.eol_uow.persist_report(session, generated)
    return generated


def _all_docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_pass_and_fail_reports_are_complete_readable_and_previewed(tmp_path: Path):
    with TestClient(app) as client:
        with isolated_data_services(tmp_path) as (_database, generator):
            passed = seed_session(state.database, generator, "EOL-PHASE03-PASS", "PASS")
            failed = seed_session(state.database, generator, "EOL-PHASE03-FAIL", "FAIL")
            for generated, expected in ((passed, "PASS"), (failed, "FAIL")):
                assert set(generated["files"]) == {"json", "docx", "pdf", "csv"}
                assert all(status == "COMPLETED" for status in generated["statuses"].values())
                payload = json.loads(Path(generated["files"]["json"]).read_text(encoding="utf-8"))
                assert payload["summary"]["result"] == expected
                assert len(payload["steps"]) == 3
                assert len(payload["assertions"]) == 3
                trace = payload["traceability"]
                assert trace["software_version"] == "v1.0.2"
                assert trace["dbc_hash"] == "dbc-phase03-hash"
                assert trace["config_hash"] == "config-phase03-hash"
                assert trace["test_plan_id"] == "standard-eol"
                assert trace["test_plan_version"] == "1.0.2"
                assert trace["operator"] == "op-phase03"

                docx_text = _all_docx_text(Path(generated["files"]["docx"]))
                assert "低速无人车" in docx_text
                assert "BMS检测" in docx_text
                with fitz.open(generated["files"]["pdf"]) as pdf:
                    pdf_text = "\n".join(page.get_text() for page in pdf)
                    assert "低速无人车" in pdf_text
                    assert "BMS检测" in pdf_text
                    assert "□" not in pdf_text
                with Path(generated["files"]["csv"]).open("r", encoding="utf-8-sig", newline="") as stream:
                    csv_rows = list(csv.DictReader(stream))
                assert len(csv_rows) == 3
                assert {row["result"] for row in csv_rows} == {expected}

            pdf_id = next(item for item in passed["database_ids"] if item.endswith("-pdf"))
            preview = client.get(f"/api/v1/reports/{pdf_id}/preview")
            assert preview.status_code == 200
            body = preview.json()
            assert body["file_sha256"] == passed["hashes"]["pdf"]
            assert body["preview_image_data_url"].startswith("data:image/png;base64,")
            assert any("低速无人车" in line for line in body["document_text"])


def test_pdf_report_summarizes_oversized_assertion_values_without_layout_failure(tmp_path: Path):
    generator = ReportGenerator(tmp_path / "reports")
    session = {
        "id": "EOL-PHASE03-LONG-VALUE",
        "chassis_no": "YL-LONG-001",
        "vin": "LLONG000000000001",
        "serial_no": "SN-LONG-001",
        "operator": "op-phase03",
        "station_id": "EOL-STATION-01",
        "status": "FAILED",
        "overall_result": "FAIL",
        "failure_reason": "反馈采样窗口不满足断言",
    }
    measured = [
        {"expected": 30.0, "actual": float(index % 3), "error": 30.0 - float(index % 3), "sample": index}
        for index in range(400)
    ]
    steps = [{
        "id": "steering_check",
        "order": 8,
        "name": "转向检测",
        "status": "FAILED",
        "result": "FAIL",
        "duration_ms": 2000,
        "failure_reason": "转向反馈未跟随",
        "assertions": [{
            "assertion_id": "front_steering_follow",
            "description": "前转角命令与反馈一致",
            "signal_name": "SAS_Front_Angle",
            "threshold": {"max_error": 8.0},
            "measured_value": measured,
            "unit": "deg",
            "quality": "good",
            "source_can_id": "0xE1",
            "source_channel": "CAN1",
            "result": "FAIL",
        }],
    }]
    generated = generator.generate(
        session,
        steps,
        metadata={"software_version": "v1.0.2", "dbc_hash": "dbc", "config_hash": "cfg", "test_plan_version": "2.0.0"},
    )
    assert generated["statuses"]["pdf"] == "COMPLETED"
    with fitz.open(generated["files"]["pdf"]) as document:
        text = "\n".join(page.get_text() for page in document)
    assert "完整值见 JSON/CSV 报告" in text
    payload = json.loads(Path(generated["files"]["json"]).read_text(encoding="utf-8"))
    assert len(payload["assertions"][0]["measured_value"]) == 400


def test_report_history_file_actions_permissions_and_traversal(tmp_path: Path, auth_headers):
    with TestClient(app) as client:
        with isolated_data_services(tmp_path) as (database, generator):
            passed = seed_session(database, generator, "EOL-PHASE03-PASS", "PASS")
            failed = seed_session(database, generator, "EOL-PHASE03-FAIL", "FAIL")
            report_id = next(item for item in failed["database_ids"] if item.endswith("-pdf"))
            report_path = Path(database.query_one("SELECT file_path FROM reports WHERE id=?", (report_id,))["file_path"])

            denied = client.delete(f"/api/v1/reports/{report_id}", headers=auth_headers("operator"))
            forged = client.delete(f"/api/v1/reports/{report_id}", headers={**auth_headers("operator"), "x-role": "admin"})
            assert denied.status_code == forged.status_code == 403
            deleted = client.delete(f"/api/v1/reports/{report_id}", headers=auth_headers("admin"))
            assert deleted.status_code == 200
            assert not report_path.exists()
            assert database.query_one("SELECT id FROM reports WHERE id=?", (report_id,)) is None
            audit = database.query_one("SELECT * FROM operator_actions WHERE action_type='delete_report' ORDER BY id DESC LIMIT 1")
            assert audit and audit["role"] == "admin" and audit["trace_id"]

            outside = tmp_path / "outside.json"
            outside.write_text("{}", encoding="utf-8")
            database.execute(
                "INSERT INTO reports(id,session_id,chassis_no,vin,result,report_type,file_path,file_size_bytes,generation_status,generated_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                ("outside-report", "EOL-PHASE03-PASS", "YL-PASS-001", "LPASS00000000001", "PASS", "json", str(outside), 2, "COMPLETED", _utc(10)),
            )
            traversal = client.get("/api/v1/reports/outside-report/preview")
            assert traversal.status_code == 403
            assert traversal.json()["code"] == "PATH_OUTSIDE_ALLOWED_ROOT"
            database.execute("DELETE FROM reports WHERE id='outside-report'")
            missing = client.get("/api/v1/reports/does-not-exist/preview")
            assert missing.status_code == 404

            history = client.get("/api/v1/history/dashboard", params={"result": "PASS", "page": 1, "page_size": 1}).json()
            assert history["pagination"]["total"] == 1
            assert history["sessions"][0]["session_id"] == "EOL-PHASE03-PASS"
            empty = client.get("/api/v1/history/dashboard", params={"vin": "NO-MATCH"}).json()
            assert empty["sessions"] == [] and empty["selected_session"] is None

            session_id = "EOL-PHASE03-PASS"
            assert len(client.get(f"/api/v1/test-sessions/{session_id}/timeline").json()) == 3
            assert len(client.get(f"/api/v1/test-sessions/{session_id}/operator-actions").json()) >= 1
            downloads = client.get(f"/api/v1/test-sessions/{session_id}/downloads").json()
            assert len(downloads) == 5 and all(item["available"] for item in downloads)
            for file_type in ("raw-can", "decoded-signals", "report-bundle", "audit-log", "curve-replay"):
                response = client.get(f"/api/v1/test-sessions/{session_id}/download/{file_type}", headers=auth_headers("viewer"))
                assert response.status_code == 200, (file_type, response.text)
                assert len(response.content) > 0

            replay = client.get(f"/api/v1/test-sessions/{session_id}/replay").json()
            assert replay["point_count"] == 2
            assert replay["mock"] is False
            historical_curve = client.get(
                "/api/v1/signals/timeseries",
                params={"mode": "history", "session_id": session_id},
            ).json()
            voltage = next(series for series in historical_curve["charts"]["bms"]["series"] if series["name"] == "BMS_Total_Voltage")
            assert voltage["data"] == [329.6, 330.1]

            exported = client.post("/api/v1/history/export", json={"result": "PASS"}, headers=auth_headers("operator"))
            assert exported.status_code == 200
            export_url = exported.json()["details"]["download_url"]
            export_file = client.get(f"/api/v1{export_url}", headers=auth_headers("viewer"))
            assert export_file.status_code == 200
            assert b"EOL-PHASE03-PASS" in export_file.content

            pareto = client.get("/api/v1/history/statistics").json()["charts"]["failure_pareto"]
            assert pareto["cumulative_percent"][-1] == 100.0

            remaining_pdf = next(item for item in passed["database_ids"] if item.endswith("-pdf"))
            print_job = client.post(f"/api/v1/reports/{remaining_pdf}/print", json={}, headers=auth_headers("operator")).json()["details"]
            persisted = database.query_one("SELECT * FROM report_print_jobs WHERE id=?", (print_job["job_id"],))
            assert persisted and persisted["status"] == "QUEUED"
